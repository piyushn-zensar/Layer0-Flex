"""J-11: the CSV and Excel compliance matrix. Compliance reads Fully / Partially / Non-compliant; Reason, Conditions,
Workaround and Next Steps are appended last (blank until A-18 stores them); everything else stays where it was.

Run:  .venv\\Scripts\\python -m pytest -q tests/test_compliance_exports.py
Uses a temporary store, so it never touches your local database (no model call: frozen answers only).
"""
import csv
import io
import os
import tempfile
from types import SimpleNamespace

os.environ["STORE_DIR"] = tempfile.mkdtemp()
os.environ["LLM_PROVIDER"] = "mock"

from conftest import signed_in  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.core.db import SessionLocal  # noqa: E402
from app.modules.consolidation import excel, service  # noqa: E402
from app.modules.workpackages.models import Assignment  # noqa: E402
from scripts import seed_demo  # noqa: E402

ORIGINAL = ["Requirement ID", "Source", "Category", "Requirement", "Quote", "Business unit", "Product", "Offering type",
            "Compliance", "Response", "Validated by", "State", "Assignment status", "Responded by"]
NEW = ["Reason", "Conditions", "Workaround", "Next Steps"]
CUSTOMER_ORIGINAL = ["Requirement ID", "RFP reference", "Requirement", "RFP wording", "Compliance", "Response", "Offered by"]


def _seeded():
    seed_demo.main()  # 13-item Syracuse seed: dispatched, some answers submitted or validated, none partial or non-compliant
    client = signed_in()
    return client, client.get("/api/portfolio").json()[0]["opp"]["id"]


def _csv(client, opp) -> list[list[str]]:
    text = client.get(f"/api/bids/{opp}/compliance-matrix.csv").text
    assert text.startswith("﻿")  # still UTF-8 with a byte order mark, so Excel reads it as UTF-8
    return list(csv.reader(io.StringIO(text.lstrip("﻿"))))


def _xlsx(client, opp):
    r = client.get(f"/api/bids/{opp}/compliance-matrix.xlsx")
    assert r.status_code == 200
    return load_workbook(io.BytesIO(r.content))


def _fake_req(req_id="REQ-9-0001", **kw):
    return SimpleNamespace(**({"req_id": req_id, "source": "p. 1, line 2", "category": "technical", "text": "Text", "quote": "Quote",
                               "kind": "item", "document_id": "d"} | kw))


def _fake_answer(**kw):
    return SimpleNamespace(**({"id": 1, "bu": "CROWN", "status": "validated", "compliance": "partial", "response": "Lead time 14 weeks",
                               "product_ref": "CROWN-ARMV", "validated_by": "Bid Manager", "responded_by": "Crown PM"} | kw))


def test_csv_keeps_its_columns_and_appends_the_four_new_ones():
    client, opp = _seeded()
    table = _csv(client, opp)
    head, body = table[0], table[1:]
    assert head == ORIGINAL + NEW and service.COLUMNS == head  # originals in their old positions, new ones last
    assert all(len(r) == 18 for r in body)
    assert all(r[14:] == ["", "", "", ""] for r in body)  # A-18 values unavailable: blank cells, never a placeholder
    assert "Not provided" not in "".join("".join(r) for r in table)
    ids = [r[0] for r in body]
    assert ids == [r[0] for r in _csv(client, opp)[1:]]  # deterministic order: a second download is the same
    assert len(body) == 13  # one row per assignment (13 requirements, one unit each)


def test_csv_compliance_wording_and_one_row_per_assignment():
    client, opp = _seeded()
    first = opp.replace("OPP-", "REQ-") + "-0001"  # CROWN, validated, met
    with SessionLocal() as db:
        db.add(Assignment(opportunity_id=opp, req_id=first, bu="EP2", owner="EP² Product Manager", status="submitted",
                          compliance="partial", response="Auxiliary supply only", responded_by="EP² PM"))
        db.add(Assignment(opportunity_id=opp, req_id=opp.replace("OPP-", "REQ-") + "-0008", bu="EP2", owner="EP² Product Manager",
                          status="submitted", compliance="not_met", response="Not offered", responded_by="EP² PM"))
        db.commit()
    body = _csv(client, opp)[1:]
    assert len(body) == 15  # two more assignments, two more rows
    mine = {r[5]: r for r in body if r[0] == first}
    assert set(mine) == {"CROWN", "EP2"}  # one requirement, two units: each keeps its own answer
    assert mine["CROWN"][8] == "Fully" and mine["EP2"][8] == "Partially" and mine["EP2"][9] == "Auxiliary supply only"
    words = {r[8] for r in body}
    assert {"Fully", "Partially", "Non-compliant", "Exception", ""} >= words and "Non-compliant" in words  # "Exception" is the seed's earlier value
    assert not words & {"met", "partial", "not_met", "Comply", "Partially comply", "Does not comply"}
    assert next(r for r in body if r[8] == "") [10] == ""  # a row with no result stays blank (assigned: no validation either)
    # the stored values are untouched: the matrix endpoint still reports them
    stored = {(r["req_id"], r["unit"]): r["compliance"] for r in client.get(f"/api/bids/{opp}/compliance-matrix").json()["rows"]}
    assert stored[(first, "EP2")] == "partial"


def test_csv_extended_cells_are_filled_and_formula_safe():
    req, a = _fake_req(text='=HYPERLINK("http://x")'), _fake_answer(reason="=1+1", conditions="Needs site survey", workaround="+cmd",
                                                                  next_steps="Ask the customer")
    row = service._row(req, None, a, "pending")
    assert len(row) == 18 and row[3] == "'=HYPERLINK(\"http://x\")"  # the existing guard still applies
    assert row[8] == "Partially"
    assert row[14:] == ["'=1+1", "Needs site survey", "'+cmd", "Ask the customer"]  # new cells are guarded too
    assert service._row(req, None, None, "not_assigned")[14:] == ["", "", "", ""]  # a requirement nobody holds
    assert service._row(req, None, SimpleNamespace(**{k: v for k, v in vars(a).items() if k not in ("reason", "workaround")}), "pending")[14:] \
        == ["", "Needs site survey", "", "Ask the customer"]  # a field missing from the object: blank


def test_excel_sheets_headers_and_formatting():
    client, opp = _seeded()
    wb = _xlsx(client, opp)
    assert wb.sheetnames == ["Compliance matrix", "Tracking"]
    ws, tr = wb["Compliance matrix"], wb["Tracking"]
    assert ws["A1"].value.startswith("Compliance matrix: ") and ws["A1"].font.bold and ws["A2"].value.startswith("Customer: ")
    head = [c.value for c in ws[4] if c.value]
    assert head == CUSTOMER_ORIGINAL + NEW  # customer sheet: originals first, new ones appended
    assert [c.value for c in tr[1] if c.value] == ORIGINAL + NEW
    assert ws.freeze_panes == "A5" and tr.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A4:K{ws.max_row}" and tr.auto_filter.ref == f"A1:R{tr.max_row}"
    for sheet, first in ((ws, 5), (tr, 2)):
        assert sheet.max_row > first
        assert all(c.alignment.wrap_text for c in sheet[first])  # wrapped text kept on data rows
        assert all(c.alignment.wrap_text for c in sheet[first - 1] if c.value)  # and on the header
    widths = {name: ws.column_dimensions[chr(65 + i)].width for i, name in enumerate(head)}
    assert widths["Requirement"] == 48 and widths["Response"] == 60  # existing widths unchanged
    assert all(widths[n] >= 40 for n in NEW)  # room for long future values
    assert all(tr.column_dimensions[chr(65 + i)].width >= 40 for i in range(14, 18))
    assert all(cell.data_type != "f" for sheet in wb for row in sheet.iter_rows() for cell in row)  # never a formula


def test_excel_wording_blank_extended_cells_and_deterministic_cells():
    client, opp = _seeded()
    wb = _xlsx(client, opp)
    ws, tr = wb["Compliance matrix"], wb["Tracking"]
    words = {row[4] for row in ws.iter_rows(min_row=5, values_only=True)}
    assert "Fully" in words and words <= {"Fully", "Partially", "Non-compliant", "Exception", "Open"}
    assert {row[8] for row in tr.iter_rows(min_row=2, values_only=True)} <= {"Fully", "Partially", "Non-compliant", "Exception", None, ""}
    for sheet, start, cols in ((ws, 5, range(7, 11)), (tr, 2, range(14, 18))):
        assert all(row[i] in (None, "") for row in sheet.iter_rows(min_row=start, values_only=True) for i in cols)  # blank until A-18
    again = _xlsx(client, opp)
    for name in wb.sheetnames:
        a = [[c.value for c in r] for r in wb[name].iter_rows(min_row=3 if name != "Tracking" else 1)]
        b = [[c.value for c in r] for r in again[name].iter_rows(min_row=3 if name != "Tracking" else 1)]
        assert a == b  # the same cells in the same order on every download (the prepared date sits in rows 1-2)


def test_excel_customer_sheet_with_populated_extended_fields_and_two_units():
    """Test-only objects stand in for A-18's fields. Only validated answers reach the customer; one line per unit."""
    crown = _fake_answer(id=1, bu="CROWN", compliance="partial", reason="Lead time", conditions="Site survey", workaround="Phase 1 now",
                         next_steps="Confirm date", response="14 weeks")
    ep2 = _fake_answer(id=2, bu="EP2", compliance="met", response="Fits", product_ref=None)
    open_one = _fake_answer(id=3, bu="CROWN", status="submitted", compliance="not_met", reason="Not yet validated", response="No")
    rows = [{"req": _fake_req("REQ-9-0001"), "assignments": [crown, ep2]},
            {"req": _fake_req("REQ-9-0002"), "assignments": [open_one]},
            {"req": _fake_req("REQ-9-0003"), "assignments": [_fake_answer(id=4, compliance="not_met", reason="=SUM(A1)")]}]
    tracking = [service._row(r["req"], None, a, "answered") for r in rows for a in r["assignments"]]
    wb = load_workbook(io.BytesIO(excel.build("Test bid", "Customer", rows, {}, lambda c: c.title(), service.COLUMNS, tracking)))
    ws = wb["Compliance matrix"]
    data = {r[0]: r for r in ws.iter_rows(min_row=5, values_only=True)}
    first = data["REQ-9-0001"]
    assert first[4] == "Partially"  # the worst validated answer wins
    assert first[7:] == ("Crown: Lead time", "Crown: Site survey", "Crown: Phase 1 now", "Crown: Confirm date")  # EP2 has none: no line
    assert data["REQ-9-0002"][4] == "Open" and data["REQ-9-0002"][7:] == (None, None, None, None)  # not validated: nothing shown
    assert data["REQ-9-0003"][4] == "Non-compliant" and data["REQ-9-0003"][7] == "Crown: =SUM(A1)"
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)  # stays text, never a formula
    track = list(wb["Tracking"].iter_rows(min_row=2, values_only=True))
    assert track[0][14:] == ("Lead time", "Site survey", "Phase 1 now", "Confirm date") and track[1][14:] == (None, None, None, None)
    assert track[2][8] == "Non-compliant" and track[3][14] == "=SUM(A1)"  # Tracking shows the text itself, as typed


def test_unknown_opportunity_is_still_a_404_for_both_exports():
    client, _ = _seeded()
    assert client.get("/api/bids/OPP-9999/compliance-matrix.csv").status_code == 404
    assert client.get("/api/bids/OPP-9999/compliance-matrix.xlsx").status_code == 404


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main(["-q", __file__]))
