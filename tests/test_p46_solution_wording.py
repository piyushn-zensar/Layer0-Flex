"""P-46 (C-40): screens and exports show "Proposed solution: unit · solution area", never a catalog id or part number;
the hand-off keeps the family id for the unit's systems.

Run:  .venv\\Scripts\\python -m pytest -q tests/test_p46_solution_wording.py
Uses a temporary store, so it never touches your local database (no model call).
"""
import csv
import io
import os
import tempfile

os.environ["STORE_DIR"] = tempfile.mkdtemp()
os.environ["LLM_PROVIDER"] = "mock"

from conftest import signed_in  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.modules.catalog import service as catalog  # noqa: E402
from scripts import seed_demo  # noqa: E402


def test_solution_helper():
    assert catalog.solution("CROWN", "CROWN-ARMV") == {"unit": "Crown Technical Systems", "area": "Arc-resistant metal-clad MV switchgear",
                                                       "examples": ["Crown arc-resistant metal-clad switchgear"]}
    own = catalog.solution("CROWN", "Two-section lineup")  # a unit's own text, not a catalog family
    assert own["area"] == "Two-section lineup" and own["examples"] == []
    assert catalog.solution("CROWN", None)["area"] is None
    assert len(catalog.solution("FPM", "FPM-DCDC")["examples"]) == 3


def test_screens_and_exports_never_show_a_catalog_id():
    seed_demo.main()
    client = signed_in()
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    ids = {p["id"] for p in catalog.products()}

    rows = client.get(f"/api/bids/{opp}/trace").json()["rows"]
    matched = [r for r in rows if r["match"] and r["match"]["bu"]]
    assert matched and all(r["solutions"] for r in matched)
    assert all(s["area"] not in ids for r in rows for s in r["solutions"])

    items = [i for i in client.get("/api/inbox/CROWN").json()["items"] if i["product_ref"]]
    assert items and all(i["solution"]["area"] not in ids for i in items)

    matrix = client.get(f"/api/bids/{opp}/compliance-matrix").json()["rows"]
    assert any(r["solution"] for r in matrix) and all(r["solution"]["area"] not in ids for r in matrix if r["solution"])

    table = list(csv.reader(io.StringIO(client.get(f"/api/bids/{opp}/compliance-matrix.csv").content.decode("utf-8-sig"))))
    col = table[0].index("Proposed solution")
    cells = [r[col] for r in table[1:]]
    assert any(" · " in c for c in cells) and not any(c in ids or "(CROWN-" in c for c in cells)

    wb = load_workbook(io.BytesIO(client.get(f"/api/bids/{opp}/compliance-matrix.xlsx").content))
    assert wb["Compliance matrix"].cell(row=4, column=7).value == "Proposed solution"


def test_handoff_keeps_the_family_id():
    seed_demo.main()
    client = signed_in()
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    payload = client.get(f"/api/bids/{opp}/handoff/CROWN").json()
    items = [i for route in payload["routes"].values() for i in route["items"]]
    assert any(i["product_id"] for i in items)
    for i in items:  # a unit manager's non-technical item has no solution (product_id None)
        assert i["proposed_solution"] == (catalog.solution("CROWN", i["product_id"]) if i["product_id"] else None)
        if i["product_id"]:
            assert i["proposed_solution"]["area"] == catalog.product(i["product_id"])["name"]
