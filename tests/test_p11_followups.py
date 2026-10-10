"""P-11 follow-ups: exports name the addendum a requirement is anchored in; a group's highlight follows its current
sub-requirements after a change document revised or removed one of them.

Run:  .venv\\Scripts\\python -m pytest -q tests/test_p11_followups.py
Uses a temporary store, so it never touches your local database (no model call: frozen answers only).
"""
import io
import json
import os
import tempfile
from pathlib import Path

os.environ["STORE_DIR"] = tempfile.mkdtemp()
os.environ["LLM_PROVIDER"] = "mock"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402
from scripts import seed_demo  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
ADDENDUM = ROOT / "data/RFP/samples/rfp_syracuse_addendum_1.pdf"
BM = "Bid Manager"


def test_exports_name_the_addendum():
    """The illustrative Addendum No. 1 on the 13-item seed: the revised breaker line (REQ-....-0005) and the added
    requirements are anchored in the addendum, so every export says so; main-RFP sources read as before."""
    seed_demo.main()  # its own opportunity: frozen, go, dispatched, some answers validated
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    s = client.post(f"/api/opportunities/{opp}/changes",
                    files={"file": (ADDENDUM.name, ADDENDUM.read_bytes(), "application/pdf")}).json()
    client.post(f"/api/changes/{s['id']}/confirm-all")
    res = client.post(f"/api/changes/{s['id']}/apply").json()["result"]
    breaker = opp.replace("OPP-", "REQ-") + "-0005"
    assert breaker in res["modified"]
    rows = {r["req"]["req_id"]: r["req"] for r in client.get(f"/api/opportunities/{opp}/trace").json()["rows"]}
    assert rows[breaker]["document_id"] == s["document_id"] and rows[breaker]["source"] == "p. 1, lines 14-15"
    named = f"{ADDENDUM.name}, p. 1, lines 14-15"
    changed = {k: f"{ADDENDUM.name}, {rows[k]['source']}" for k in res["modified"] + res["added"]}
    main = {k: r["source"] for k, r in rows.items() if k not in changed}
    assert main and all(v.startswith("p. ") for v in main.values())

    # CSV: same columns, the Source names the addendum
    import csv as _csv
    text = client.get(f"/api/opportunities/{opp}/compliance-matrix.csv").text.lstrip("﻿")
    table = list(_csv.reader(io.StringIO(text)))
    head, body = table[0], table[1:]
    assert head[:2] == ["Requirement ID", "Source"] and len(head) == 14
    sources = {r[0]: r[1] for r in body}
    assert sources[breaker] == named and all(sources[k] == v for k, v in changed.items())
    assert all(sources[k] == v for k, v in main.items())  # unchanged for the main RFP
    assert "p. 55, line 16" in sources.values()

    # Excel: the customer sheet's RFP reference and the Tracking sheet's Source
    wb = load_workbook(io.BytesIO(client.get(f"/api/opportunities/{opp}/compliance-matrix.xlsx").content))
    customer = {r[0]: r[1] for r in wb["Compliance matrix"].iter_rows(min_row=5, values_only=True) if r[0]}
    tracking = {r[0]: r[1] for r in wb["Tracking"].iter_rows(min_row=2, values_only=True)}
    for got in (customer, tracking):
        assert got[breaker] == named and all(got[k] == v for k, v in changed.items())
        assert all(got[k] == v for k, v in main.items())

    # response outline: JSON "source" fields and the Markdown
    outline = client.get(f"/api/opportunities/{opp}/response-outline").json()
    refs = {x["req_id"]: x["source"] for c in outline["chapters"] for x in c["open"] + c["material"]}
    assert refs[breaker] == named and all(refs[k] == v for k, v in changed.items() if k in refs)
    assert all(refs[k] == v for k, v in main.items() if k in refs)
    md = client.get(f"/api/opportunities/{opp}/response-outline.md").text
    assert f"{breaker} ({named})" in md


def test_group_highlight_follows_current_sub_requirements():
    """A group made of three page-55 lines; the addendum revises one (now in the addendum) and removes another. The
    group's highlight on the RFP page keeps only the remaining line; the stored versions are not touched; the Excel
    customer sheet names the addendum for the revised sub-requirement."""
    from app.core.db import SessionLocal
    from app.modules.ingestion import service as ingestion
    from app.modules.opportunities import service as opportunities
    from app.modules.requirements import service as requirements

    seed = json.loads((ROOT / "data/seed/demo_syracuse.json").read_text("utf-8"))
    oracle = json.loads((ROOT / "data/seed/syracuse_addendum_1_expected.json").read_text("utf-8"))
    with SessionLocal() as db:
        opp = opportunities.create(db, "group highlight", "", "", BM).id
        pdf = ROOT / seed["document"]
        doc = opportunities.add_document(db, opp, pdf.name, pdf.read_bytes(), "main", BM)
        ingestion.ingest(db, doc.id)
        items = [requirements.add(db, opp, doc.id, i["quote"], i["text"], i["category"], "test", i["section"], i["page"])
                 for i in seed["requirements"]]
        autocad, breaker, arc = items[3], items[4], items[5]  # p. 55, lines 39-40, 46-49, 52-53
        assert {autocad.page, breaker.page, arc.page} == {55} and all(r.bboxes for r in (autocad, breaker, arc))
        group = requirements._create_group(db, opp, [autocad, breaker, arc], "Switchgear lineup", "technical", "test")
        db.commit()
        ids = {"group": group.req_id, "autocad": autocad.req_id, "breaker": breaker.req_id, "arc": arc.req_id}
        old = {k: [tuple(b) for b in r.bboxes] for k, r in (("autocad", autocad), ("breaker", breaker), ("arc", arc))}
        for r in requirements.current(db, opp):
            requirements.review(db, r.req_id, "approve", BM)
        requirements.freeze(db, opp, BM)
        change = opportunities.add_document(db, opp, ADDENDUM.name, ADDENDUM.read_bytes(), "change", BM)
        ingestion.ingest(db, change.id)
        quote = next(o["quote"] for o in oracle["items"]
                     if (o["baseline_quote"] or "").startswith("Metal Clad Switchgear shall have"))
        removed = next(o["quote"] for o in oracle["items"] if o["kind"] == "removed")
        out = requirements.apply_change(db, opp, change.id, [
            {"kind": "modified", "target": ids["breaker"], "quote": quote, "category": "technical", "page": 1,
             "text": "Twenty (20) 15kV vacuum circuit breaker positions."},
            {"kind": "removed", "target": ids["autocad"], "quote": removed}], BM, "test addendum")
        assert out["modified"] == [ids["breaker"]] and out["removed"] == [ids["autocad"]] and out["affected"] == [ids["group"]]

    client = TestClient(app)
    trace = client.get(f"/api/opportunities/{opp}/trace").json()
    rows = {r["req"]["req_id"]: r for r in trace["rows"]}
    g = rows[ids["group"]]
    boxes = {tuple(b) for b in g["req"]["bboxes"]}
    assert boxes == set(old["arc"])  # only the line still in force on this page
    assert not boxes & (set(old["breaker"]) | set(old["autocad"]))
    assert [k["req_id"] for k in g["children"]] == [ids["arc"], ids["breaker"]]  # the RFP first, then the addendum
    assert next(k for k in g["children"] if k["req_id"] == ids["breaker"])["document_id"] == change.id
    assert g["req"]["page"] == 55 and "bboxes" not in g  # the field keeps its place and name
    item = next(r for r in trace["rows"] if r["req"]["kind"] == "item")
    stored = {r["req_id"]: r for r in client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]}
    assert item["req"]["bboxes"] == stored[item["req"]["req_id"]]["bboxes"]  # items as stored
    assert {tuple(b) for b in stored[ids["group"]]["bboxes"]} >= set(old["breaker"]) | set(old["autocad"])  # immutable

    wb = load_workbook(io.BytesIO(client.get(f"/api/opportunities/{opp}/compliance-matrix.xlsx").content))
    customer = {r[0]: r[1] for r in wb["Compliance matrix"].iter_rows(min_row=5, values_only=True) if r[0]}
    assert customer[ids["breaker"]] == f"{ADDENDUM.name}, p. 1, lines 14-15"
    assert customer[ids["arc"]] == "p. 55, lines 52-53" and ids["autocad"] not in customer
