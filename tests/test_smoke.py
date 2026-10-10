"""Smoke test: the demo flow runs through every module and every screen renders.

Run:  .venv\\Scripts\\python -m pytest -q
Uses a temporary store, so it never touches your local database.
"""
import os
import re
import tempfile
from pathlib import Path

os.environ["STORE_DIR"] = tempfile.mkdtemp()
os.environ["LLM_PROVIDER"] = "mock"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import text  # noqa: E402

from app.core.db import engine  # noqa: E402
from app.main import app  # noqa: E402
from scripts import seed_demo  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]


def standalone(rows: list[dict]) -> list[dict]:
    """Requirements that are not groups, sub-requirements or duplicates (the reader's own line items)."""
    return [r for r in rows if r["kind"] == "item" and not r["parent_id"] and r["status"] == "proposed"
            and r["provenance"] == "EXTRACTED"]


def same_page(rows: list[dict], n: int) -> list[dict]:
    """The first page with at least n stand-alone requirements."""
    from collections import defaultdict
    pages = defaultdict(list)
    for r in standalone(rows):
        pages[r["page"]].append(r)
    return next(v for _, v in sorted(pages.items()) if len(v) >= n)


def test_demo_flow_and_api():
    seed_demo.main()
    client = TestClient(app)
    for url in ["/api/portfolio", "/api/people", "/api/opportunities/OPP-0001",
                "/api/opportunities/OPP-0001/requirements", "/api/opportunities/OPP-0001/trace",
                "/api/opportunities/OPP-0001/decisions", "/api/opportunities/OPP-0001/consolidation",
                "/api/opportunities/OPP-0001/changes", "/api/inbox/CROWN", "/api/inbox/EP2", "/api/inbox/BID",
                "/api/catalog?q=relay"]:
        assert client.get(url).status_code == 200, url
    trace = client.get("/api/opportunities/OPP-0001/trace").json()
    by_id = {r["req"]["req_id"]: r for r in trace["rows"]}
    assert by_id["REQ-0001-0006"]["req"]["source"].startswith("p. 55")
    assert by_id["REQ-0001-0002"]["unit"]["code"] == "EP2"
    assert client.get(f"/api/documents/{trace['doc']['id']}/pages/55.png").headers["content-type"] == "image/png"
    csv = client.get("/api/opportunities/OPP-0001/compliance-matrix.csv").text
    assert "p. 55, lines" in csv

    # rule R6: the same file in a second opportunity is that opportunity's own document
    opp2 = client.post("/api/opportunities", json={"title": "second"}).json()["id"]
    pdf = (ROOT / "data/RFP/RFP-2023-20-Switchgear-Procurement-Final.pdf").read_bytes()
    doc2 = client.post(f"/api/opportunities/{opp2}/documents", files={"file": ("rfp.pdf", pdf)}).json()
    assert doc2["opportunity_id"] == opp2 and doc2["id"] != trace["doc"]["id"]
    assert client.get(f"/api/opportunities/{opp2}").json()["documents"][0]["status"] == "ingested"

    # actor header is URL-encoded by the web app (EP² is not ASCII)
    a = next(i for i in client.get("/api/inbox/EP2").json()["items"] if not i["responded_by"])
    r = client.post(f"/api/assignments/{a['id']}/respond", json={"compliance": "met", "response": "ok"},
                    headers={"X-Actor": "EP%C2%B2%20Design%20Engineer"})
    assert r.status_code == 200
    after = {i["id"]: i for i in client.get("/api/inbox/EP2").json()["items"]}
    assert after[a["id"]]["responded_by"] == "EP² Design Engineer"

    with pytest.raises(Exception, match="append-only"):  # rule R8
        with engine.begin() as conn:
            conn.execute(text("UPDATE audit_event SET actor = 'someone else'"))


def test_modules_talk_only_through_services():
    """Modular-monolith rule: module A may import module B's service.py, never its models/agent/controller."""
    for f in (ROOT / "app" / "modules").rglob("*.py"):
        own = f.relative_to(ROOT / "app" / "modules").parts[0]
        for mod, part in re.findall(r"app\.modules\.(\w+)\.(\w+)", f.read_text("utf-8")):
            assert mod == own or part == "service", f"{f.name} imports app.modules.{mod}.{part}"


def test_people_decide_and_dispatch_needs_go():
    """Rule R4: re-running matching keeps people's decisions; dispatch needs a go; bad input is refused.
    Appended by Atharv (append-only on this shared file)."""
    seed_demo.main()  # its own opportunity, so this test does not depend on test order
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]  # newest first
    rows = client.get(f"/api/opportunities/{opp}/trace").json()["rows"]
    rejected, accepted = rows[0]["match"], rows[1]["match"]
    assert client.post(f"/api/matches/{rejected['id']}/decide", json={"action": "reject"}).status_code == 200
    assert client.post(f"/api/matches/{accepted['id']}/decide", json={"action": "accept"}).status_code == 200
    assert client.post(f"/api/opportunities/{opp}/match").json()["kept"] == 2
    after = {r["req"]["req_id"]: r["match"] for r in client.get(f"/api/opportunities/{opp}/trace").json()["rows"]}
    assert after[rejected["req_id"]] is None  # stays rejected: no older proposal comes back
    assert after[accepted["req_id"]]["id"] == accepted["id"] and after[accepted["req_id"]]["status"] == "accepted"

    a = client.get("/api/inbox/CROWN").json()["items"][0]
    assert client.post(f"/api/matches/{accepted['id']}/decide", json={"action": "maybe"}).status_code == 422
    assert client.post("/api/matches/999999/decide", json={"action": "accept"}).status_code == 404
    assert client.post(f"/api/assignments/{a['id']}/respond", json={"compliance": "yes"}).status_code == 422
    assert client.post("/api/assignments/999999/respond", json={"compliance": "met"}).status_code == 404
    assert client.post(f"/api/opportunities/{opp}/go-no-go", json={"outcome": "maybe"}).status_code == 422
    assert client.get(f"/api/opportunities/{opp}/decisions").json()["go_no_go"]["outcome"] == "go"

    client.post(f"/api/opportunities/{opp}/go-no-go", json={"outcome": "no_go"})
    assert client.post(f"/api/opportunities/{opp}/dispatch").status_code == 409
    assert client.get(f"/api/opportunities/{opp}").json()["opportunity"]["status"] == "no_go"


def test_requirement_review_actions():
    """P-06: edit, split, merge, add missed; freeze needs every line item decided; replaced items leave the flow."""
    client = TestClient(app)
    opp = client.post("/api/opportunities", json={"title": "review actions"}).json()["id"]
    pdf = (ROOT / "data/RFP/RFP-2023-20-Switchgear-Procurement-Final.pdf").read_bytes()
    client.post(f"/api/opportunities/{opp}/documents", files={"file": ("rfp.pdf", pdf)})
    assert client.post(f"/api/opportunities/{opp}/requirements/extract").json()["proposed"] > 100  # frozen GPT-4o answers
    reqs = {r["req_id"]: r for r in client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]}
    p55 = same_page(list(reqs.values()), 4)  # four stand-alone requirements on one page

    edited = client.post(f"/api/requirements/{p55[0]['req_id']}/review",
                         json={"action": "edit", "text": "Edited text", "category": "schedule"}).json()
    assert (edited["text"], edited["category"]) == ("Edited text", "schedule")
    assert client.post(f"/api/requirements/{p55[0]['req_id']}/review", json={"action": "edit", "category": "nope"}).status_code == 409

    q = p55[1]["quote"]
    half = q.rfind(" ", 0, len(q) // 2)
    parts = client.post(f"/api/requirements/{p55[1]['req_id']}/split",
                        json={"parts": [{"quote": q[:half]}, {"quote": q[half + 1:]}]}).json()
    assert len(parts) == 2 and all(p["provenance"] == "EXTRACTED" and p["derived_from"] == [p55[1]["req_id"]] for p in parts)

    merged = client.post(f"/api/opportunities/{opp}/requirements/merge",
                         json={"req_ids": [p55[2]["req_id"], p55[3]["req_id"]], "text": "One obligation"}).json()
    assert merged["derived_from"] == [p55[2]["req_id"], p55[3]["req_id"]] and merged["provenance"] == "EXTRACTED"
    assert client.post(f"/api/requirements/{p55[2]['req_id']}/split",
                       json={"parts": [{"quote": "a"}, {"quote": "b"}]}).status_code == 409  # already merged

    missed = client.post(f"/api/opportunities/{opp}/requirements",
                         json={"quote": "Each phase shall have 1 inch diameter ground ball.", "category": "technical", "page": 55}).json()
    assert missed["provenance"] == "EXTRACTED" and missed["page"] == 55

    assert client.post(f"/api/opportunities/{opp}/baselines").status_code == 409  # undecided items left
    rows = client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]
    for r in rows:
        if r["status"] == "proposed":
            client.post(f"/api/requirements/{r['req_id']}/review", json={"action": "approve"})
    assert client.post(f"/api/opportunities/{opp}/baselines").status_code == 200
    statuses = {r["req_id"]: r["status"] for r in client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]}
    assert statuses[p55[1]["req_id"]] == "split" and statuses[p55[2]["req_id"]] == "merged"
    trace_ids = {r["req"]["req_id"] for r in client.get(f"/api/opportunities/{opp}/trace").json()["rows"]}
    assert p55[1]["req_id"] not in trace_ids and merged["req_id"] in trace_ids  # replaced items leave the flow


def test_one_requirement_several_units(monkeypatch):
    """A-04: the matcher may list several units; dispatch creates one assignment per participating unit.
    The matcher is stubbed (no model call). Appended by Atharv (append-only on this shared file)."""
    from app.core.db import SessionLocal
    from app.modules.decisions import service as decisions
    from app.modules.matching import agent
    from app.modules.matching import service as matching
    from app.modules.workpackages import service as workpackages

    def two_units(header, requirements, products, past):  # per requirement: two catalog products from different units
        units = {}
        for p in products:
            units.setdefault(p["bu"], {"bu": p["bu"], "product_id": p["id"], "offering_type": p["offering_type"]})
        pick = list(units.values())[:2]
        return [{"n": i, "units": pick, "confidence": 0.9, "rationale": "stub"} for i in range(1, len(requirements) + 1)]

    monkeypatch.setattr(agent, "propose_page", two_units)  # P-18: one matcher call per page
    seed_demo.main()  # its own opportunity (matched retrieval-only, then dispatched)
    with SessionLocal() as db:
        opp = TestClient(app).get("/api/portfolio").json()[0]["opp"]["id"]
        matching.match(db, opp, "test")  # proposals are refreshed with the stub's two units
        multi = {r: m for r, m in matching.for_opportunity(db, opp).items() if len({u["bu"] for u in m.units}) == 2}
        assert multi, "stub should give some requirement two units"
        req_id, m = next(iter(multi.items()))
        assert (m.bu, m.product_id) == (m.units[0]["bu"], m.units[0]["product_id"])  # main unit first
        both = [u["bu"] for u in m.units]
        assert all(req_id in matching.suggested_units(db, opp)[bu] for bu in both)
        decisions.record(db, opp, "participation", "units", list(matching.suggested_units(db, opp)), "", "test")
        decisions.record(db, opp, "go_no_go", "go", [], "", "test")
        workpackages.dispatch(db, opp, "test")
        assert {a.bu for a in workpackages.by_requirement(db, opp)[req_id]} >= set(both)
        assert workpackages.dispatch(db, opp, "test") == 0  # repeatable: nothing duplicated


def test_person_changes_match():
    """A-05: a person accepts, rejects or changes a match (several units); re-dispatch adds the new unit only.
    Appended by Atharv (append-only on this shared file)."""
    seed_demo.main()  # its own opportunity, dispatched
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    rows = client.get(f"/api/opportunities/{opp}/trace").json()["rows"]
    req_id = next(r["req"]["req_id"] for r in rows if r["match"] and r["match"]["bu"] == "CROWN")
    url = f"/api/opportunities/{opp}/requirements/{req_id}/match"
    body = {"units": [{"product_id": "CROWN-ARMV", "offering_type": "ETO"}, {"product_id": "EP2-RPP", "offering_type": "ETO"}]}
    m = client.post(url, json=body, headers={"X-Actor": "Bid%20Manager"}).json()
    assert (m["status"], m["method"], m["decided_by"], m["bu"]) == ("accepted", "manual", "Bid Manager", "CROWN")
    assert [u["bu"] for u in m["units"]] == ["CROWN", "EP2"]  # unit follows the product
    assert client.post(url, json={"units": [{"product_id": "NOPE", "offering_type": "ETO"}]}).status_code == 422
    assert client.post(url, json={"units": [{"product_id": "EPC-800V", "offering_type": "ETO"}]}).status_code == 422  # pending unit
    assert client.post(url, json={"units": [{"product_id": "EP2-RPP", "offering_type": "NONE"}]}).status_code == 422
    assert client.post(f"/api/opportunities/{opp}/requirements/REQ-9999-0001/match", json=body).status_code == 404

    client.post(f"/api/opportunities/{opp}/participation", json={"units": ["CROWN", "EP2"]})
    client.post(f"/api/opportunities/{opp}/go-no-go", json={"outcome": "go"})
    assert client.post(f"/api/opportunities/{opp}/dispatch").json()["created"] >= 1  # the added EP2 work
    row = next(r for r in client.get(f"/api/opportunities/{opp}/trace").json()["rows"] if r["req"]["req_id"] == req_id)
    assert sorted({a["bu"] for a in row["assignments"]}) == ["CROWN", "EP2"]

    none = client.post(url, json={"units": []}).json()  # a person says: not a product item
    assert none["bu"] is None and none["offering_type"] == "NONE" and none["units"] == []


def test_workflow_order_and_roles():
    """Fixes from the 8 Oct system test: decide and dispatch only after freeze; only the unit answers and only
    the Bid Manager validates submitted work; re-dispatch withdraws work that no longer fits.
    Appended by Atharv (append-only on this shared file)."""
    client = TestClient(app)
    BM, CROWN = {"X-Actor": "Bid%20Manager"}, {"X-Actor": "Crown%20Design%20Engineer"}
    fresh = client.post("/api/opportunities", json={"title": "not frozen"}).json()["id"]
    assert client.post(f"/api/opportunities/{fresh}/go-no-go", json={"outcome": "go"}).status_code == 409
    assert client.post(f"/api/opportunities/{fresh}/dispatch").status_code == 409
    assert client.post("/api/opportunities/OPP-9999/participation", json={"units": ["CROWN"]}).status_code == 404
    for bad in ([], ["XYZ"], ["EPC"]):  # empty, unknown, pending
        assert client.post(f"/api/opportunities/{fresh}/participation", json={"units": bad}).status_code == 422

    seed_demo.main()
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    item = next(i for i in client.get("/api/inbox/CROWN").json()["items"] if i["opportunity_id"] == opp and i["status"] == "assigned")
    url = f"/api/assignments/{item['id']}"
    assert client.post(f"{url}/validate", json={"ok": True}, headers=BM).status_code == 409        # nothing submitted
    assert client.post(f"{url}/respond", json={"compliance": "met"}, headers={"X-Actor": "EP%C2%B2%20Design%20Engineer"}).status_code == 403
    assert client.post(f"{url}/respond", json={"compliance": "met"}, headers=CROWN).status_code == 200
    assert client.post(f"{url}/respond", json={"compliance": "met"}, headers=CROWN).status_code == 409  # already submitted
    assert client.post(f"{url}/validate", json={"ok": True}, headers=CROWN).status_code == 403      # no self-validation
    assert client.post(f"{url}/validate", json={"ok": True}, headers=BM).status_code == 200

    # the bid manager takes the item away from Crown: re-dispatch withdraws Crown's work and gives it to the bid desk
    assert client.post(f"/api/opportunities/{opp}/requirements/{item['req_id']}/match", json={"units": []}, headers=BM).status_code == 200
    client.post(f"/api/opportunities/{opp}/dispatch", headers=BM)
    row = next(r for r in client.get(f"/api/opportunities/{opp}/trace").json()["rows"] if r["req"]["req_id"] == item["req_id"])
    assert [a["bu"] for a in row["assignments"]] == ["BID"]
    assert all(i["id"] != item["id"] for i in client.get("/api/inbox/CROWN").json()["items"])     # gone from Crown's inbox
    assert client.get("/api/catalog/search?q=relay&k=-1").status_code == 422


def test_review_findings_fixed():
    """Review of 9 Oct: path traversal, freeze guards, unknown opportunities, parallel create, ID reuse, input checks."""
    import concurrent.futures

    client = TestClient(app)
    pdf = (ROOT / "data/RFP/RFP-2023-20-Switchgear-Procurement-Final.pdf").read_bytes()
    sha = "893403680eda487a7fcf4968231102e10345324aefd0503d2b28b97d3347a76e"

    # page endpoint: only real documents and real pages (no file paths, no page 0)
    outside = str(ROOT / "data" / "layout_cache" / sha).replace("/", "\\")
    assert client.get(f"/api/documents/{outside}/pages/1").status_code == 404
    opp = client.post("/api/opportunities", json={"title": "guards"}).json()["id"]
    doc = client.post(f"/api/opportunities/{opp}/documents", files={"file": ("rfp.pdf", pdf)}).json()
    assert "path" not in doc
    assert client.get(f"/api/documents/{doc['id']}/pages/0").status_code == 404
    assert client.get(f"/api/documents/{doc['id']}/pages/102.png").status_code == 404
    assert client.post(f"/api/opportunities/{opp}/documents", data={"role": "main"},
                       files={"file": ("again.pdf", pdf + b"\n")}).status_code == 409  # one main RFP

    # unknown opportunity: 404 and nothing stored
    assert client.post("/api/opportunities/OPP-9996/documents", files={"file": ("x.pdf", pdf)}).status_code == 404
    assert client.post("/api/opportunities/OPP-9997/baselines").status_code == 404

    # freeze: needs approved items; once only; nothing changes afterwards
    assert client.post(f"/api/opportunities/{opp}/baselines").status_code == 409  # nothing read or approved yet
    client.post(f"/api/opportunities/{opp}/requirements/extract")
    everything = client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]
    first_ids = {r["req_id"] for r in everything}
    rows = [r for r in everything if not r["parent_id"] and r["status"] == "proposed"]  # requirements to decide
    client.post(f"/api/requirements/{rows[0]['req_id']}/review", json={"action": "reject"})
    assert client.post(f"/api/opportunities/{opp}/requirements/extract").status_code == 409  # review started
    assert client.post(f"/api/opportunities/{opp}/requirements", json={"quote": "  ", "category": "technical"}).status_code == 409
    assert client.post(f"/api/opportunities/{opp}/requirements/merge",
                       json={"req_ids": [rows[1]["req_id"], rows[2]["req_id"]], "text": " "}).status_code == 409
    alone = standalone(everything)
    p55 = alone[0]
    p56 = next(r for r in alone if r["page"] != p55["page"])
    assert client.post(f"/api/opportunities/{opp}/requirements/merge",
                       json={"req_ids": [p55["req_id"], p56["req_id"]], "text": "x"}).status_code == 409  # cross-page
    for r in rows[1:]:
        client.post(f"/api/requirements/{r['req_id']}/review", json={"action": "approve"})
    assert client.post(f"/api/opportunities/{opp}/baselines").status_code == 200
    assert client.post(f"/api/opportunities/{opp}/baselines").status_code == 409  # no empty baseline 2
    assert client.post(f"/api/requirements/{rows[0]['req_id']}/review", json={"action": "approve"}).status_code == 409
    assert len(client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]) == len(first_ids)

    # status never moves backwards
    client.post(f"/api/opportunities/{opp}/participation", json={"units": ["CROWN"]})
    client.post(f"/api/opportunities/{opp}/go-no-go", json={"outcome": "go"})
    client.post(f"/api/opportunities/{opp}/dispatch")
    client.post(f"/api/opportunities/{opp}/go-no-go", json={"outcome": "go"})
    assert client.get(f"/api/opportunities/{opp}").json()["opportunity"]["status"] == "dispatched"

    # re-reading drafts never reuses an ID
    opp2 = client.post("/api/opportunities", json={"title": "reread"}).json()["id"]
    client.post(f"/api/opportunities/{opp2}/documents", files={"file": ("rfp.pdf", pdf)})
    one = {r["req_id"] for r in client.get(f"/api/opportunities/{opp2}/requirements").json()["requirements"]} or None
    client.post(f"/api/opportunities/{opp2}/requirements/extract")
    a = {r["req_id"] for r in client.get(f"/api/opportunities/{opp2}/requirements").json()["requirements"]}
    assert client.post(f"/api/opportunities/{opp2}/requirements/extract").status_code == 200  # nothing reviewed yet
    b = {r["req_id"] for r in client.get(f"/api/opportunities/{opp2}/requirements").json()["requirements"]}
    assert len(a) == len(b) and not (a & b) and one is None

    # parallel creates all succeed with distinct IDs
    with concurrent.futures.ThreadPoolExecutor(10) as pool:
        made = list(pool.map(lambda i: client.post("/api/opportunities", json={"title": f"p{i}"}), range(10)))
    assert all(r.status_code == 200 for r in made) and len({r.json()["id"] for r in made}) == 10

    # acting user is cleaned and capped
    long = client.post("/api/opportunities", json={"title": "actor"}, headers={"X-Actor": "A" * 5000 + "%0D%0A"}).json()
    assert len(long["created_by"]) == 80


def test_trace_and_consolidation_findings_fixed():
    """Review of 9 Oct (Janvia's modules): CSV formula cells, unknown opportunity 404, no evidence blob in trace."""
    client = TestClient(app)
    assert client.get("/api/opportunities/OPP-9999/trace").status_code == 404
    assert client.get("/api/opportunities/OPP-9999/consolidation").status_code == 404
    assert client.get("/api/opportunities/OPP-9999/compliance-matrix.csv").status_code == 404

    pdf = (ROOT / "data/RFP/RFP-2023-20-Switchgear-Procurement-Final.pdf").read_bytes()
    opp = client.post("/api/opportunities", json={"title": "csv"}).json()["id"]
    client.post(f"/api/opportunities/{opp}/documents", files={"file": ("rfp.pdf", pdf)})
    client.post(f"/api/opportunities/{opp}/requirements", json={
        "quote": "Switchgear shall be Arc-resistant Type 2B.", "text": "=HYPERLINK(\"http://x\")", "category": "technical", "page": 55})
    csv = client.get(f"/api/opportunities/{opp}/compliance-matrix.csv")
    assert "'=HYPERLINK" in csv.text and ',=HYPERLINK' not in csv.text
    assert csv.headers["content-disposition"] == f'attachment; filename="{opp}-compliance-matrix.csv"'
    rows = client.get(f"/api/opportunities/OPP-0001/trace").json()["rows"]
    assert rows and all(r["match"] is None or "evidence" not in r["match"] for r in rows)


def test_requirement_history():
    """P-13: every version (who, when, wording) and every event is shown; merged originals point to the merge."""
    client = TestClient(app)
    pdf = (ROOT / "data/RFP/RFP-2023-20-Switchgear-Procurement-Final.pdf").read_bytes()
    opp = client.post("/api/opportunities", json={"title": "history"}).json()["id"]
    client.post(f"/api/opportunities/{opp}/documents", files={"file": ("rfp.pdf", pdf)})
    client.post(f"/api/opportunities/{opp}/requirements/extract")
    rows = same_page(client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"], 3)
    a, b, c = (r["req_id"] for r in rows[:3])
    original = rows[0]["text"]
    new_cat = "legal" if rows[0]["category"] != "legal" else "staffing"
    client.post(f"/api/requirements/{a}/review", json={"action": "edit", "text": "Edited once", "reason": "clearer"},
                headers={"X-Actor": "Crown%20Design%20Engineer"})
    client.post(f"/api/requirements/{a}/review", json={"action": "edit", "category": new_cat})
    client.post(f"/api/requirements/{a}/review", json={"action": "approve"})
    h = client.get(f"/api/requirements/{a}/history").json()
    assert [(v["label"], v["text"], v["category"]) for v in h["versions"]] == [
        ("original", original, rows[0]["category"]), ("edited", "Edited once", rows[0]["category"]), ("edited", "Edited once", new_cat)]
    assert h["versions"][1]["by"] == "Crown Design Engineer" and h["versions"][1]["reason"] == "clearer"
    assert [e["action"] for e in h["events"]] == ["proposed", "edit", "edit", "approve"]
    merged = client.post(f"/api/opportunities/{opp}/requirements/merge", json={"req_ids": [b, c], "text": "one"}).json()["req_id"]
    assert client.get(f"/api/requirements/{b}/history").json()["events"][-1]["details"]["into"] == merged
    assert client.get(f"/api/requirements/{merged}/history").json()["derived_from"] == [b, c]
    assert client.get("/api/requirements/REQ-9999-0001/history").status_code == 404


def test_grouping_and_duplicates():
    """P-16: related line items become one requirement with sub-requirements; duplicates point to the first one."""
    client = TestClient(app)
    pdf = (ROOT / "data/RFP/RFP-2023-20-Switchgear-Procurement-Final.pdf").read_bytes()
    opp = client.post("/api/opportunities", json={"title": "grouping"}).json()["id"]
    client.post(f"/api/opportunities/{opp}/documents", files={"file": ("rfp.pdf", pdf)})
    r = client.post(f"/api/opportunities/{opp}/requirements/extract").json()  # frozen reader + grouping answers
    assert r["proposed"] == 814 and r["duplicates"] == 21 and r["groups"] > 100 and r["requirements"] < 400, r
    rows = client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]
    by_id = {x["req_id"]: x for x in rows}
    group = next(x for x in rows if x["kind"] == "group")
    kids = [x for x in rows if x["parent_id"] == group["req_id"]]
    assert len(kids) >= 2 and all(k["page"] == group["page"] for k in kids)
    assert {tuple(b) for k in kids for b in k["bboxes"]} <= {tuple(b) for b in group["bboxes"]}  # all highlights kept
    dup = next(x for x in rows if x["status"] == "duplicate")
    assert by_id[dup["derived_from"][0]]["status"] != "duplicate"

    # the workflow sees requirements, not sub-requirements
    trace = client.get(f"/api/opportunities/{opp}/trace").json()["rows"]
    assert len(trace) == r["requirements"] and all(t["req"]["parent_id"] is None for t in trace)
    assert next(t for t in trace if t["req"]["req_id"] == group["req_id"])["children"]

    # approving a group decides its sub-requirements; ungroup releases them
    client.post(f"/api/requirements/{kids[0]['req_id']}/review", json={"action": "reject"})
    client.post(f"/api/requirements/{group['req_id']}/review", json={"action": "approve"})
    after = {x["req_id"]: x["status"] for x in client.get(f"/api/opportunities/{opp}/requirements").json()["requirements"]}
    assert after[kids[0]["req_id"]] == "rejected" and all(after[k["req_id"]] == "approved" for k in kids[1:])
    other = next(x for x in rows if x["kind"] == "group" and x["req_id"] != group["req_id"])
    released = client.post(f"/api/requirements/{other['req_id']}/ungroup").json()
    assert released and all(x["parent_id"] is None for x in released)
    assert client.post(f"/api/requirements/{other['req_id']}/split", json={"parts": [{"quote": "a"}, {"quote": "b"}]}).status_code == 409
    assert client.post(f"/api/requirements/{dup['req_id']}/review", json={"action": "approve"}).json()["status"] == "approved"


def test_returned_items_reopen_with_a_note():
    """A-07: the Bid Manager returns an answer only with a note (design 7.5); the unit sees the note and answers
    the returned item again. Appended by Atharv (append-only on this shared file)."""
    seed_demo.main()
    client = TestClient(app)
    BM, CROWN = {"X-Actor": "Bid%20Manager"}, {"X-Actor": "Crown%20Design%20Engineer"}
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    item = next(i for i in client.get("/api/inbox/CROWN").json()["items"] if i["opportunity_id"] == opp and i["status"] == "assigned")
    url = f"/api/assignments/{item['id']}"
    client.post(f"{url}/respond", json={"compliance": "partial", "response": "first try"}, headers=CROWN)
    assert client.post(f"{url}/validate", json={"ok": False, "note": "  "}, headers=BM).status_code == 422
    assert client.post(f"{url}/validate", json={"ok": False, "note": "Name the breaker model."}, headers=BM).status_code == 200
    back = next(i for i in client.get("/api/inbox/CROWN").json()["items"] if i["id"] == item["id"])
    assert (back["status"], back["validation_note"], back["response"]) == ("returned", "Name the breaker model.", "first try")
    assert client.post(f"{url}/respond", json={"compliance": "met", "response": "VCB type X"}, headers=CROWN).status_code == 200
    assert next(i for i in client.get("/api/inbox/CROWN").json()["items"] if i["id"] == item["id"])["status"] == "submitted"


def test_go_no_go_summary_and_structured_criteria():
    """A-12: summary of how requirements are satisfied, criteria with 'data not yet available' placeholders,
    advice that never decides, and the person's judgement per criterion stored with the decision.
    Appended by Atharv (append-only on this shared file)."""
    seed_demo.main()
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    s = client.get(f"/api/opportunities/{opp}/decisions").json()["summary"]
    assert s["advice"].startswith("Advice from Layer 0") and s["advice"].endswith("A person decides.")
    status = {c["id"]: c["status"] for c in s["criteria"]}
    assert all(status[k] == "unknown" for k in ("cost_budget", "delivery", "competitors"))  # deviations, capacity: A-10
    assert status["open_questions"] == "not_met"  # Syracuse states no redundancy class (R-003 warn)
    assert {r["level"] for r in s["rows"]} <= {"fully", "partly", "not", "bid_desk"} and len(s["rows"]) == 13
    crit = [{"id": "coverage", "status": "met", "note": "checked"}, {"id": "cost_budget", "status": "unknown"}]
    d = client.post(f"/api/opportunities/{opp}/go-no-go", json={"outcome": "go", "criteria": crit}, headers={"X-Actor": "Bid%20Manager"}).json()
    assert d["criteria"][0] == {"id": "coverage", "status": "met", "note": "checked"} and "summary" in d["evidence"]
    assert client.post(f"/api/opportunities/{opp}/go-no-go", json={"outcome": "go", "criteria": [{"id": "x", "status": "met"}]}).status_code == 422


def test_unit_handoff_payload():
    """A-09 (M8): per-unit hand-off JSON grouped by route; every payload needs a person; nothing claims a design.
    Appended by Atharv (append-only on this shared file)."""
    seed_demo.main()
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    r = client.get(f"/api/opportunities/{opp}/handoff/CROWN")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    p = r.json()
    assert p["requires_human_completion"] is True and p["unit"]["code"] == "CROWN"
    items = [i for route in p["routes"].values() for i in route["items"]]
    assert items and all(i["quote"] and i["source"] for i in items)
    for i in p["routes"].get("basis_of_design", {}).get("items", []):
        assert i["basis_of_design"]["spec_claimed"] is False
    for i in p["routes"].get("cpq_seed", {}).get("items", []):
        assert i["offering_type"] == "CTO" and "bom_lines" in i["configuration_seed"]
    assert {i["offering_type"] for i in items} >= {"CTO", "ETO"}  # Crown accessories (CTO) and switchgear (ETO)
    assert client.get(f"/api/opportunities/{opp}/handoff/BID").status_code == 409
    assert client.get("/api/opportunities/OPP-9999/handoff/CROWN").status_code == 404


def test_compliance_matrix_new_columns_and_bom():
    """J-05 part C: "Assignment status" and "Responded by" are appended (earlier columns keep their place), and the file
    starts with a UTF-8 byte order mark so Excel reads non-ASCII text. Appended by Janvia (append-only on this shared file)."""
    import csv
    import io

    seed_demo.main()
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    r = client.get(f"/api/opportunities/{opp}/compliance-matrix.csv")
    assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    head = rows[0]
    assert head[:12] == ["Requirement ID", "Source", "Category", "Requirement", "Quote", "Business unit", "Product",
                         "Offering type", "Compliance", "Response", "Validated by", "State"]
    assert head[12:] == ["Assignment status", "Responded by"] and all(len(row) == len(head) for row in rows)
    status, by = head.index("Assignment status"), head.index("Responded by")
    assert {row[status] for row in rows[1:]} >= {"validated", "assigned"}
    assert all(row[by] for row in rows[1:] if row[status] in ("submitted", "validated"))  # an answer always has a person


def test_rfp_search_short_term_index():
    """The opportunity's RFP is indexed when it is loaded (short-term RAG) and searchable with exact sources."""
    client = TestClient(app)
    pdf = (ROOT / "data/RFP/RFP-2023-20-Switchgear-Procurement-Final.pdf").read_bytes()
    opp = client.post("/api/opportunities", json={"title": "rfp search"}).json()["id"]
    assert client.get(f"/api/opportunities/{opp}/rfp-search", params={"q": "insurance"}).status_code == 404  # not read yet
    doc = client.post(f"/api/opportunities/{opp}/documents", files={"file": ("rfp.pdf", pdf)}).json()
    found = client.get(f"/api/opportunities/{opp}/rfp-search", params={"q": "general liability insurance limits", "k": 3}).json()
    assert found["passages"] and found["passages"][0]["page"] == 13, found  # the insurance requirements
    p = found["passages"][0]
    assert p["line_start"] <= p["line_end"] and "insurance" in p["text"].lower()
    assert client.get("/api/opportunities/OPP-9999/rfp-search", params={"q": "x"}).status_code == 404


def test_compliance_matrix_excel():
    """J-08: customer sheet (validated answers only, customer words, sub-requirements listed) and internal tracking."""
    import io

    from openpyxl import load_workbook

    seed_demo.main()  # its own opportunity: 13 line items, some answers validated
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    r = client.get(f"/api/opportunities/{opp}/compliance-matrix.xlsx")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('-compliance-matrix.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Compliance matrix", "Tracking"]
    ws = wb["Compliance matrix"]
    assert [c.value for c in ws[4]][:5] == ["Requirement ID", "RFP reference", "Requirement", "RFP wording", "Compliance"]
    words = {row[4] for row in ws.iter_rows(min_row=5, values_only=True)}
    assert "Comply" in words and words <= {"Comply", "Partially comply", "Does not comply", "Exception", "Open"}
    assert all(cell.data_type != "f" for sheet in wb for row in sheet.iter_rows() for cell in row)  # never a formula
    assert wb["Tracking"].max_row - 1 == len(client.get(f"/api/opportunities/{opp}/compliance-matrix.csv").text.strip().splitlines()) - 1
    assert client.get("/api/opportunities/OPP-9999/compliance-matrix.xlsx").status_code == 404


def test_bid_and_portfolio_checks():
    """A-10: data-sheet ratings vs the standard products in play, and unit workload vs capacity, as go/no-go evidence."""
    seed_demo.main()  # its own opportunity: matched to CROWN and EP2 products, dispatched
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    d = client.get(f"/api/opportunities/{opp}/decisions").json()
    dev, load = d["evidence"]["deviations"], d["evidence"]["workload"]
    assert dev["placeholder"] and load["placeholder"]  # never presented as confirmed figures
    assert "CROWN-ARMV" in dev["products"] and dev["checked"] >= 1
    current = next(r for r in dev["rows"] if r["field"] == "Rated continuous current")  # data sheet asks 2000/1200 A
    assert current["severity"] == "high" and current["source"].startswith("p. 74")
    crit = {c["id"]: c for c in d["summary"]["criteria"]}
    assert crit["deviations"]["status"] == "not_met" and "p. 74" in crit["deviations"]["detail"]
    units = {r["bu"]: r for r in load["rows"]}
    assert {"CROWN", "EP2"} <= set(units) and all(r["total"] >= r["open_here"] for r in units.values())
    assert crit["capacity"]["status"] == "met" and "placeholder" in crit["capacity"]["detail"]


def test_response_outline_agent():
    """P-10 + J-07: the drafting agent writes only from validated answers (frozen drafts, no model call), every
    paragraph cites them, open requirements are listed and never drafted, and the outline downloads as Markdown."""
    seed_demo.main()  # its own opportunity: 13 line items, 4 validated answers
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    o = client.get(f"/api/opportunities/{opp}/response-outline").json()
    validated = {a["req_id"] for c in o["chapters"] for a in c["material"]}
    assert o["validated_answers"] == len(validated) > 0 and o["summary"]["drafted"]
    for d in [o["summary"]] + [c["draft"] for c in o["chapters"]]:
        assert all(p["sources"] and set(p["sources"]) <= validated for p in d["paragraphs"])  # cites validated answers only
    by_id = {c["id"]: c for c in o["chapters"]}
    assert by_id["technical"]["draft"]["drafted"] and sum(c["total"] for c in o["chapters"]) == o["total"]
    assert all(not c["draft"]["drafted"] for c in o["chapters"] if not c["material"])  # nothing drafted without answers
    assert all(r["req_id"] not in validated or r["state"] != "answered" for c in o["chapters"] for r in c["open"])
    md = client.get(f"/api/opportunities/{opp}/response-outline.md")
    assert md.status_code == 200 and md.text.startswith("# Response outline:") and "## 1. Technical response" in md.text
    assert client.get("/api/opportunities/OPP-9999/response-outline").status_code == 404


def test_knowledge_base_queue():
    """A-11: people send a requirement, a validated answer or a decision rationale; only the curator approves;
    approved items join the long-term search index (never the matcher's fixed prompt part)."""
    from app.modules.catalog import service as catalog

    seed_demo.main()  # its own opportunity, some answers validated
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    unit = {"X-Actor": "Crown%20Product%20Manager"}
    items = [a for rows in [client.get("/api/inbox/CROWN").json()["items"]] for a in rows if a["opportunity_id"] == opp]
    done = next(a for a in items if a["status"] == "validated")
    open_ = next(a for a in items if a["status"] != "validated")
    past_before = catalog.past_responses()

    kb = client.post("/api/knowledge", json={"kind": "response", "ref": str(done["id"]), "note": "reuse"}, headers=unit).json()
    assert kb["status"] == "queued" and kb["sent_by"] == "Crown Product Manager" and kb["response"].startswith("Met.")
    assert client.post("/api/knowledge", json={"kind": "response", "ref": str(done["id"])}).status_code == 409  # already queued
    assert client.post("/api/knowledge", json={"kind": "response", "ref": str(open_["id"])}).status_code == 409  # not validated
    assert client.post("/api/knowledge", json={"kind": "requirement", "ref": "REQ-9999-0001"}).status_code == 404
    go = client.get(f"/api/opportunities/{opp}/decisions").json()["go_no_go"]
    dec = client.post("/api/knowledge", json={"kind": "decision", "ref": str(go["id"])}).json()
    assert dec["response"].startswith("go:")

    review = f"/api/knowledge/{kb['kb_id']}/review"
    assert client.post(review, json={"approve": True}, headers=unit).status_code == 403  # only the curator
    assert client.post(f"/api/knowledge/{dec['kb_id']}/review", json={"approve": False}).status_code == 409  # reject needs a note
    ok = client.post(review, json={"approve": True, "response": "Met. Two-section lineup, reusable wording."}).json()
    assert ok["status"] == "approved" and ok["reviewed_by"] == "Bid Manager"
    assert client.post(review, json={"approve": True}).status_code == 409  # decided once
    assert client.post(f"/api/knowledge/{dec['kb_id']}/review", json={"approve": False, "note": "not general"}).json()["status"] == "rejected"

    hits = client.get("/api/catalog", params={"q": "Two-section lineup reusable wording"}).json()["results"]
    assert hits[0]["id"] == kb["kb_id"] and hits[0]["learned"]  # found by the long-term index
    assert catalog.past_responses() == past_before  # the matcher's fixed prompt part is unchanged
    sent = client.get(f"/api/opportunities/{opp}/knowledge").json()
    assert sent[f"response:{done['id']}"] == "approved" and sent[f"decision:{go['id']}"] == "rejected"
    queue = client.get("/api/knowledge", params={"status": "approved"}).json()
    assert queue["learned"] >= 1 and all(i["status"] == "approved" for i in queue["items"])


def test_change_document_against_baseline(monkeypatch, tmp_path):
    """P-11: an addendum is read and classified against the frozen baseline (model answers stubbed, no model call),
    a person confirms every item, and applying it makes baseline 2: a new version anchored in the addendum, a
    removal, a new ID; the changed answer goes back to its unit and the new requirement is dispatched."""
    from app.core.db import SessionLocal
    from app.modules.changes import agent
    from app.modules.workpackages import service as workpackages
    from scripts.txt_to_pdf import convert

    seed_demo.main()  # its own opportunity: 13 frozen line items, go, dispatched, some answers validated
    client = TestClient(app)
    BM, CROWN = {"X-Actor": "Bid%20Manager"}, {"X-Actor": "Crown%20Design%20Engineer"}
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    rows = client.get(f"/api/opportunities/{opp}/trace").json()["rows"]
    by_quote = lambda start: next(r for r in rows if r["req"]["quote"].startswith(start))
    arc, ball = by_quote("Switchgear shall be Arc-resistant Type 2B"), by_quote("Each phase shall have 1 inch")
    done = next(a for a in arc["assignments"] if a["status"] == "validated")
    ball_work = [a["id"] for a in ball["assignments"]]
    assert ball_work

    lines = ["ADDENDUM NO. 1", "Bidders shall acknowledge receipt of this addendum on the bid form.",
             "Switchgear shall be Arc-resistant Type 2C (replaces Type 2B).",
             "Delete: Each phase shall have 1 inch diameter ground ball.",
             "Add: The switchgear shall include one spare 1200 A vacuum breaker.",
             "Q1: Is the Equipment Building by others? A1: Yes, no change."]
    (tmp_path / "addendum-1.txt").write_text("\n".join(lines), "utf-8")
    pdf = convert(tmp_path / "addendum-1.txt").read_bytes()
    upload = lambda: client.post(f"/api/opportunities/{opp}/changes", files={"file": ("addendum-1.pdf", pdf)})

    r = upload()  # mock mode, no frozen answer: refused, no change set
    assert r.status_code == 409 and "frozen answers" in r.json()["detail"]
    assert client.get(f"/api/opportunities/{opp}/changes").json()["sets"] == []

    said = [("info", "submission", lines[1], "Acknowledge the addendum."),
            ("modify", "compliance", "L3: " + lines[2], "Switchgear shall be arc-resistant Type 2C."),
            ("delete", "technical", lines[3], "The 1 inch ground ball on each phase is deleted."),
            ("add", "technical", lines[4], "Include one spare 1200 A vacuum circuit breaker."),
            ("clarify", "technical", lines[5], "The equipment building remains by others.")]
    want = {"modify": ("modified", "Arc-resistant Type 2B"), "delete": ("removed", "ground ball"),
            "clarify": ("unchanged", "Equipment Building"), "add": ("added", None)}

    def model(task, system, prompt, schema):  # stands in for the frozen answers of both tasks
        if task == "read_changes":
            assert "L3: Switchgear shall be Arc-resistant Type 2C" in prompt
            return {"statements": [{"page": 1, "quote": q, "text": t, "category": c, "action": a} for a, c, q, t in said]}
        assert "REQ-" not in prompt and "acknowledge" not in prompt  # no IDs; info is decided by rule
        answers = []
        for block in re.split(r"^\[", prompt, flags=re.M)[1:]:
            n, action = re.match(r"(\d+)\] Reader says: (\w+)", block).groups()
            kind, words = want[action]
            letter = next((l for l, text in re.findall(r"^\s+([A-E])\. (.*)$", block, re.M) if words and words in text), "")
            answers.append({"n": int(n), "change": kind, "target": letter, "new_text": "", "rationale": "stub", "confidence": 0.8})
        return {"answers": answers}

    monkeypatch.setattr(agent, "complete_json", model)
    s = upload().json()
    assert s["status"] == "review" and s["baseline_from"] == 1 and s["total"] == 5 and s["confirmed"] == 0
    items = {i["action"]: i for i in s["items"]}
    assert [i["proposed_kind"] for i in s["items"]] == ["not_a_requirement", "modified", "removed", "added", "unchanged"]
    mod = items["modify"]
    assert mod["quote"] == lines[2] and mod["source"] == "p. 1, line 3" and mod["bboxes"]  # anchored in the addendum
    assert mod["proposed_target"] == arc["req"]["req_id"] and mod["target_text"] == arc["req"]["text"]
    assert arc["req"]["req_id"] in [c["req_id"] for c in mod["candidates"]] and len(mod["candidates"]) <= 5
    assert items["delete"]["proposed_target"] == ball["req"]["req_id"] and items["add"]["proposed_target"] is None
    assert s["counts"] == {"added": 1, "modified": 1, "removed": 1, "unchanged": 1, "not_a_requirement": 1}
    assert s["share"] == round(2 / 13, 3) and s["drastic"] is False and s["pages"][0]["page"] == 1
    assert upload().status_code == 409  # one set in review at a time

    url = f"/api/changes/{s['id']}"
    q = items["clarify"]
    assert client.post(f"{url}/items/{q['id']}", json={"kind": "modified", "target": "REQ-9999-0001"}).status_code == 409
    assert client.post(f"{url}/items/{q['id']}", json={"kind": "removed"}).status_code == 409  # needs a target
    assert client.post(f"{url}/items/999999", json={"kind": "added"}).status_code == 404
    over = client.post(f"{url}/items/{q['id']}", json={"kind": "not_a_requirement", "target": arc["req"]["req_id"]},
                       headers=CROWN).json()
    assert (over["kind"], over["target"], over["status"], over["decided_by"]) == ("not_a_requirement", None, "confirmed",
                                                                                   "Crown Design Engineer")
    add = client.post(f"{url}/items/{items['add']['id']}", json={"kind": "added", "text": "One spare 1200 A breaker."}).json()
    assert add["text"] == "One spare 1200 A breaker."
    assert client.post(f"{url}/apply", headers=BM).status_code == 409    # three items not confirmed yet
    assert client.post(f"{url}/apply", headers=CROWN).status_code == 403  # only the Bid Manager
    assert client.post(f"{url}/confirm-all").json()["confirmed"] == 5

    applied = client.post(f"{url}/apply", headers=BM).json()
    res = applied["result"]
    assert applied["status"] == "applied" and applied["baseline_to"] == 2 and applied["applied_by"] == "Bid Manager"
    assert res["modified"] == [arc["req"]["req_id"]] and res["removed"] == [ball["req"]["req_id"]] and len(res["added"]) == 1
    assert res["returned"] >= 1 and res["dispatched"] >= 1 and res["note"] == ""
    assert client.post(f"{url}/apply", headers=BM).status_code == 409
    assert client.post(f"{url}/discard", headers=BM).status_code == 409

    reqs = client.get(f"/api/opportunities/{opp}/requirements").json()
    assert reqs["baseline"]["number"] == 2 and reqs["baseline"]["count"] == 13  # one removed, one added
    by_id = {r["req_id"]: r for r in reqs["requirements"]}
    new_arc, new = by_id[arc["req"]["req_id"]], by_id[res["added"][0]]
    assert (new_arc["version"], new_arc["text"], new_arc["baseline"], new_arc["status"]) == (
        2, "Switchgear shall be arc-resistant Type 2C.", 2, "approved")
    assert new_arc["document_id"] == s["document_id"] and new_arc["source"] == "p. 1, line 3"
    assert by_id[ball["req"]["req_id"]]["status"] == "removed"
    assert new["req_id"] not in {r["req"]["req_id"] for r in rows} and (new["status"], new["baseline"]) == ("approved", 2)
    assert new["text"] == "One spare 1200 A breaker." and new["document_id"] == s["document_id"]
    h = client.get(f"/api/requirements/{arc['req']['req_id']}/history").json()
    assert [v["label"] for v in h["versions"]] == ["original", "revised"] and "addendum-1.pdf" in h["versions"][1]["reason"]
    assert [e["action"] for e in h["events"]].count("frozen") == 2 and h["source"] == "p. 1, line 3"
    assert client.post(f"/api/requirements/{new['req_id']}/review", json={"action": "edit", "text": "x"}).status_code == 409

    with SessionLocal() as db:
        back = workpackages.get(db, done["id"])
        assert back.status == "returned" and back.validation_note.startswith("Changed by addendum-1.pdf: now reads:")
        assert all(workpackages.get(db, i).status == "withdrawn" for i in ball_work)
        assert workpackages.by_requirement(db, opp)[new["req_id"]]  # the added requirement is dispatched
    trace = client.get(f"/api/opportunities/{opp}/trace").json()
    assert [d["role"] for d in trace["docs"]] == ["main", "change"] and trace["docs"][1]["id"] == s["document_id"]
    assert trace["doc"]["id"] == trace["docs"][0]["id"]
    assert ball["req"]["req_id"] not in {r["req"]["req_id"] for r in trace["rows"]}
    view = client.get(f"/api/opportunities/{opp}/changes").json()
    assert view["baseline"] == {"number": 2, "count": 13} and view["sets"][0]["status"] == "applied"
    assert view["sets"][0]["result"] == res and view["drastic_threshold"] == 0.25
    assert client.get("/api/opportunities/OPP-9999/changes").status_code == 404


def test_illustrative_addendum_frozen_answers():
    """P-11 on the illustrative Addendum No. 1 with the frozen change-agent answers (no model call): every statement
    is classified as the oracle expects, and applying it gives baseline 2 without moving any untouched match."""
    import json as _json

    from app.modules.requirements.anchoring import normalise

    seed_demo.main()  # its own opportunity: the 13-item Syracuse seed, frozen, go, dispatched
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    oracle = _json.loads((ROOT / "data/seed/syracuse_addendum_1_expected.json").read_text("utf-8"))
    pdf = ROOT / oracle["document"]
    s = client.post(f"/api/opportunities/{opp}/changes", files={"file": (pdf.name, pdf.read_bytes(), "application/pdf")}).json()
    assert s["status"] == "review" and s["total"] == len(oracle["items"])
    for it in s["items"]:
        want = next(o for o in oracle["items"] if normalise(o["quote"])[:50] in normalise(it["quote"]))
        assert it["proposed_kind"] == want["kind"], (it["quote"], it["proposed_kind"])
        if want["baseline_quote"]:  # the right baseline requirement, by its RFP wording
            target = client.get(f"/api/requirements/{it['proposed_target']}/history").json()
            assert normalise(want["baseline_quote"])[:40] in normalise(target["quote"]), (it["quote"], target["quote"])
    before = {r["req"]["req_id"]: r["match"] for r in client.get(f"/api/opportunities/{opp}/trace").json()["rows"]}
    client.post(f"/api/changes/{s['id']}/confirm-all")
    done = client.post(f"/api/changes/{s['id']}/apply").json()
    assert done["status"] == "applied" and done["result"]["baseline"] == 2
    assert len(done["result"]["added"]) == 2 and len(done["result"]["modified"]) == 2 and len(done["result"]["removed"]) == 1
    after = {r["req"]["req_id"]: r["match"] for r in client.get(f"/api/opportunities/{opp}/trace").json()["rows"]}
    touched = set(done["result"]["added"] + done["result"]["modified"])
    assert all(after[k]["units"] == m["units"] for k, m in before.items() if k in after and k not in touched and m)


def test_change_set_review_findings_fixed(monkeypatch, tmp_path):
    """P-11 adversarial review: the applied set still shows the baseline wording it was compared with; two items
    that change one requirement are refused; a failure after the new baseline is committed is reported in the
    note (not a 500 with a stale result); a set taken by another session (discarded / uploaded meanwhile) is not
    applied twice; change documents keep their upload order in the requirement list."""
    from app.core.db import SessionLocal
    from app.modules.changes import agent
    from app.modules.changes import service as changes
    from app.modules.matching import service as matching
    from scripts.txt_to_pdf import convert

    seed_demo.main()  # its own opportunity: 13 frozen line items, go, dispatched, some answers validated
    client = TestClient(app)
    opp = client.get("/api/portfolio").json()[0]["opp"]["id"]
    rows = client.get(f"/api/opportunities/{opp}/trace").json()["rows"]
    breakers = next(r["req"] for r in rows if r["req"]["quote"].startswith("Metal Clad Switchgear shall have eighteen"))

    def stub(said):  # the reader returns these statements; the classifier says added or not_a_requirement
        monkeypatch.setattr(agent, "read", lambda layout: [{"page": 1, "quote": q, "text": t, "category": "technical",
                                                            "action": a} for a, q, t in said])
        monkeypatch.setattr(agent, "classify", lambda sts: [
            {"change": "added" if s["action"] == "add" else "not_a_requirement", "target": None, "new_text": s["text"],
             "rationale": "stub", "confidence": 0.5} for s in sts])

    def pdf(stem, lines):
        (tmp_path / f"{stem}.txt").write_text("\n".join(lines), "utf-8")
        return convert(tmp_path / f"{stem}.txt").read_bytes()

    one = ["ADDENDUM ONE", "Add: provide spare part number 0A.", "Replace 18 with 20 breaker positions.",
           "Replace two revenue meter positions with three."]
    stub([("add", one[1], "Spare part 0A."), ("modify", one[2], "20 breakers"), ("modify", one[3], "3 meters")])
    s = client.post(f"/api/opportunities/{opp}/changes", files={"file": ("o0.pdf", pdf("o0", one))}).json()
    url, items = f"/api/changes/{s['id']}", s["items"]
    client.post(f"{url}/items/{items[0]['id']}", json={"kind": "added"})
    client.post(f"{url}/items/{items[1]['id']}", json={"kind": "modified", "target": breakers["req_id"],
                                                        "text": "20 breaker positions, 2 revenue meter positions."})
    client.post(f"{url}/items/{items[2]['id']}", json={"kind": "modified", "target": breakers["req_id"],
                                                        "text": "18 breaker positions, 3 revenue meter positions."})
    r = client.post(f"{url}/apply")
    assert r.status_code == 409 and "all change" in r.json()["detail"]  # the first change would silently be lost
    client.post(f"{url}/items/{items[2]['id']}", json={"kind": "unchanged", "target": breakers["req_id"]})

    def down(*args, **kwargs):
        raise RuntimeError("model service unavailable")

    with monkeypatch.context() as m:
        m.setattr(matching, "match", down)
        r = client.post(f"{url}/apply")
    assert r.status_code == 200 and r.json()["status"] == "applied" and r.json()["baseline_to"] == 2
    res = r.json()["result"]
    assert "follow-up stopped" in res["note"] and res["returned"] >= 1  # the returned answer is counted
    mod = next(i for i in r.json()["items"] if i["id"] == items[1]["id"])
    assert (mod["target_text"], mod["target_source"]) == (breakers["text"], breakers["source"])  # baseline 1 wording
    crown = client.get(f"/api/opportunities/{opp}/handoff/CROWN").json()
    hand = next(i for route in crown["routes"].values() for i in route["items"] if i["req_id"] == breakers["req_id"])
    assert hand["document"] == s["document_id"] and hand["source"].startswith("p. 1")  # the addendum's page, not the RFP's

    two = ["ADDENDUM TWO", "Add: provide spare part number 0B."]
    stub([("add", two[1], "Spare part 0B.")])
    later = pdf("t0", two)

    def read_while_another_uploads(layout):  # a second upload finishes while this one is still being read
        stub([("add", two[1], "Spare part 0B.")])
        with SessionLocal() as other:
            changes.upload(other, opp, "t0.pdf", later, "Bid Manager")
        return agent.read(layout)

    monkeypatch.setattr(agent, "read", read_while_another_uploads)
    r = client.post(f"/api/opportunities/{opp}/changes", files={"file": ("r0.pdf", pdf("r0", ["RACE", "Add: x."]))})
    assert r.status_code == 409 and "still in review" in r.json()["detail"]
    pending = [x for x in client.get(f"/api/opportunities/{opp}/changes").json()["sets"] if x["status"] == "review"]
    assert [x["filename"] for x in pending] == ["t0.pdf"]

    client.post(f"/api/changes/{pending[0]['id']}/confirm-all")
    with SessionLocal() as stale:  # this session read the set while it was in review; the Bid Manager discards it
        held = changes.get(stale, pending[0]["id"])  # kept, as within one request (the identity map is weak)
        assert held.status == "review"
        assert client.post(f"/api/changes/{pending[0]['id']}/discard").json()["status"] == "discarded"
        with pytest.raises(ValueError, match="meanwhile"):
            changes.apply(stale, pending[0]["id"], "Bid Manager")
    assert client.get(f"/api/opportunities/{opp}/changes").json()["baseline"]["number"] == 2

    stub([("add", two[1], "Spare part 0B.")])
    s3 = client.post(f"/api/opportunities/{opp}/changes", files={"file": ("t0.pdf", later)}).json()
    client.post(f"/api/changes/{s3['id']}/confirm-all")
    with SessionLocal() as stale:  # a person's decision read the set before the Bid Manager applied it
        held = changes.get(stale, s3["id"])
        assert client.post(f"/api/changes/{s3['id']}/apply").json()["baseline_to"] == 3
        with pytest.raises(ValueError, match="meanwhile"):  # the record keeps what was applied
            changes.decide(stale, s3["id"], s3["items"][0]["id"], "not_a_requirement", None, None, "Crown Design Engineer")
    assert client.get(f"/api/opportunities/{opp}/changes").json()["sets"][0]["items"][0]["kind"] == "added"
    trace = client.get(f"/api/opportunities/{opp}/trace").json()
    order = [d["id"] for d in trace["docs"]]
    assert order[1:] == [s["document_id"], s3["document_id"]]
    seen = [order.index(r["req"]["document_id"]) for r in trace["rows"] if r["req"]["page"] is not None]
    assert seen == sorted(seen)  # rows follow the documents' order: main RFP, then change documents as they came in
