"""Word and Excel into the layout model (P-28 spike, C-18): the document is rendered as a plain PDF, which the normal
parser then reads, so line numbers, anchoring, highlights and page images work unchanged.

Paragraphs and table rows become lines (cells joined with " | "); each Word heading 1 and each Excel sheet starts a
new page, with the sheet name as its first line. The PDF is deterministic (fixed fonts, no dates, no new file ID), so
the same file always gives the same pages and lines. Known limits are in team/investigations.md.
"""
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pymupdf

SUFFIXES = (".docx", ".xlsx")
PAGE = pymupdf.paper_rect("letter")
MARGIN, FONT_SIZE, LEADING = 54, 10, 14
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def sections(path: str | Path) -> list[list[str]]:
    """The document as sections of lines: Word split at each heading 1, Excel one section per sheet."""
    return _docx(path) if Path(path).suffix.lower() == ".docx" else _xlsx(path)


def _docx(path) -> list[list[str]]:
    with zipfile.ZipFile(path) as z:
        body = ElementTree.fromstring(z.read("word/document.xml")).find(f"{W}body")
    out: list[list[str]] = [[]]
    for block in body:
        if block.tag == f"{W}p":
            style = block.find(f"{W}pPr/{W}pStyle")
            if style is not None and style.get(f"{W}val", "").lower() in ("heading1", "title") and out[-1]:
                out.append([])
            out[-1] += _paragraph(block)
        elif block.tag == f"{W}tbl":
            for row in block.iter(f"{W}tr"):
                cells = [" ".join(t for p in cell.iter(f"{W}p") for t in _paragraph(p)) for cell in row.iter(f"{W}tc")]
                if any(cells):
                    out[-1].append(" | ".join(c.strip() for c in cells))
    return [s for s in out if s]


def _paragraph(p) -> list[str]:
    """A paragraph's text; a manual line break starts a new line."""
    lines = [""]
    for node in p.iter():
        if node.tag == f"{W}t":
            lines[-1] += node.text or ""
        elif node.tag == f"{W}tab":
            lines[-1] += " "
        elif node.tag in (f"{W}br", f"{W}cr"):
            lines.append("")
    return [line.strip() for line in lines if line.strip()]


def _xlsx(path) -> list[list[str]]:
    from openpyxl import load_workbook  # already a dependency (compliance matrix export)

    book = load_workbook(path, read_only=True, data_only=True)  # data_only: a formula's last computed value
    try:
        out = []
        for sheet in book.worksheets:
            rows = [" | ".join("" if v is None else str(v).strip() for v in row) for row in sheet.iter_rows(values_only=True)]
            rows = [r.strip(" |") for r in rows if r.strip(" |")]
            if rows:
                out.append([f"Sheet: {sheet.title}"] + rows)
        return out
    finally:
        book.close()


def to_pdf(path: str | Path) -> bytes:
    """Render the sections as PDF pages: long lines wrap (each piece is its own line), a section starts a page."""
    return render(sections(path), Path(path).stem)[0]


def render(parts: list[list[str]], title: str) -> tuple[bytes, list[int]]:
    """Sections of lines -> deterministic PDF bytes and the page number each section starts on (its first line is
    the section's heading). Also used for web pages and text files in the knowledge base (P-35)."""
    pdf = pymupdf.open()
    width = PAGE.width - 2 * MARGIN
    starts = []
    for section in parts:
        page, y = None, 0.0
        for raw in section:
            for piece in _wrap(raw, width):
                if page is None or y > PAGE.height - MARGIN:
                    if page is None:
                        starts.append(len(pdf) + 1)
                    page, y = pdf.new_page(width=PAGE.width, height=PAGE.height), MARGIN
                page.insert_text((MARGIN, y), piece, fontname="helv", fontsize=FONT_SIZE)
                y += LEADING
    if not len(pdf):
        raise ValueError("The document has no text to read.")
    pdf.set_metadata({"title": title, "producer": "Layer 0 office.py", "creator": "Layer 0",
                      "creationDate": "", "modDate": ""})
    return pdf.tobytes(garbage=4, deflate=True, no_new_id=True), starts


def _wrap(text: str, width: float) -> list[str]:
    pieces, cur = [], ""
    for word in text.split():
        test = f"{cur} {word}".strip()
        if cur and pymupdf.get_text_length(test, fontname="helv", fontsize=FONT_SIZE) > width:
            pieces.append(cur)
            cur = word
        else:
            cur = test
    return pieces + [cur] if cur else pieces


if __name__ == "__main__":  # self-check: python -m app.modules.ingestion.office
    import tempfile

    from openpyxl import Workbook

    tmp = Path(tempfile.mkdtemp())
    book = Workbook()
    book.active.title = "Ratings"
    book.active.append(["Item", "Value", None])
    book.active.append(["Rated voltage", "15 kV", None])
    book.create_sheet("Empty")
    book.save(tmp / "t.xlsx")
    assert sections(tmp / "t.xlsx") == [["Sheet: Ratings", "Item | Value", "Rated voltage | 15 kV"]]
    xml = (f'<w:document xmlns:w="{W[1:-1]}"><w:body>'
           '<w:p><w:r><w:t>Intro line</w:t></w:r></w:p>'
           '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>2 Technical</w:t></w:r></w:p>'
           '<w:p><w:r><w:t>The switchgear shall be</w:t><w:tab/><w:t>arc-resistant.</w:t><w:br/><w:t>Next.</w:t></w:r></w:p>'
           '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Voltage</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>15 kV</w:t></w:r></w:p>'
           '</w:tc></w:tr></w:tbl><w:sectPr/></w:body></w:document>')
    with zipfile.ZipFile(tmp / "t.docx", "w") as z:
        z.writestr("word/document.xml", xml)
    assert sections(tmp / "t.docx") == [["Intro line"], ["2 Technical", "The switchgear shall be arc-resistant.", "Next.",
                                                         "Voltage | 15 kV"]]
    assert to_pdf(tmp / "t.docx") == to_pdf(tmp / "t.docx")  # deterministic: same file, same pages
    print("office ok")
