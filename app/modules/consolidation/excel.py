"""The compliance matrix as an Excel workbook (J-08). Written by code from the stored answers, never by a model.

Sheet "Compliance matrix" is what goes to the customer: one row per requirement in RFP order (a group's
sub-requirements follow it, each with its own RFP reference), the compliance in customer words, and only answers
the bid manager has validated. Sheet "Tracking" is internal: the same columns as the CSV, one row per unit answer.
Every cell is written as plain text, so nothing in it can run as a formula.
The compliance in both exports reads Fully / Partially / Non-compliant (J-11); the four extended response columns
(Reason, Conditions, Workaround, Next Steps) are appended last and stay blank until the units' answers hold them (A-18).
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY, WHITE = "0F172A", "FFFFFF"
# COMPLIANCE_WORDS is the wording the response-outline agent is prompted with (its frozen answers are keyed by that
# prompt text): leave it. EXPORT_WORDS is what the CSV and both Excel sheets show.
COMPLIANCE_WORDS = {"met": "Comply", "partial": "Partially comply", "not_met": "Does not comply", "exception": "Exception"}
EXPORT_WORDS = {"met": "Fully", "partial": "Partially", "not_met": "Non-compliant", "exception": "Exception"}
FILLS = {"Fully": "DCFCE7", "Partially": "FEF3C7", "Non-compliant": "FEE2E2", "Exception": "FEE2E2", "Open": "EEF2F6"}
WORST_FIRST = ["Non-compliant", "Exception", "Partially", "Fully"]
EXTENDED = [("reason", "Reason"), ("conditions", "Conditions"), ("workaround", "Workaround"), ("next_steps", "Next Steps")]
THIN = Side(style="thin", color="CBD5E1")
CUSTOMER_COLUMNS = [("Requirement ID", 16), ("RFP reference", 16), ("Requirement", 48), ("RFP wording", 60),
                    ("Compliance", 18), ("Response", 60), ("Offered by", 30)] + [(name, 40) for _, name in EXTENDED]
TRACKING_WIDTHS = [16, 16, 12, 40, 50, 14, 18, 12, 12, 50, 18, 14, 16, 24] + [40] * len(EXTENDED)


def customer_compliance(assignments: list) -> str:
    """One word for the customer: the worst validated answer; "Open" until every unit's answer is validated."""
    if not assignments or any(a.status != "validated" for a in assignments):
        return "Open"
    words = {EXPORT_WORDS.get(a.compliance, "Open") for a in assignments}
    return next((w for w in WORST_FIRST if w in words), "Open")


def build(title: str, customer: str, rows: list[dict], kids: dict, unit_name, tracking_columns: list[str],
          tracking_rows: list[list[str]], source=None) -> bytes:
    """`source(req)` is the RFP reference shown (the file name first for a requirement anchored in an addendum)."""
    source = source or (lambda r: r.source)
    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance matrix"
    ws.append([f"Compliance matrix: {title}"])
    ws.append([f"Customer: {customer or 'not set'}    Prepared {date.today():%d %b %Y}    "
               "Every requirement traces to its page and lines in the RFP."])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"].font = Font(italic=True, color="475569")
    header_row = 4
    _header(ws, header_row, [c for c, _ in CUSTOMER_COLUMNS])
    for row in rows:
        req, validated = row["req"], [a for a in row["assignments"] if a.status == "validated"]
        word = customer_compliance(row["assignments"])
        answer = "\n".join(f"{unit_name(a.bu)}: {a.response}" for a in validated if a.response) if word != "Open" else ""
        offered = ", ".join(dict.fromkeys(
            f"{unit_name(a.bu)} ({a.product_ref})" if a.product_ref else unit_name(a.bu) for a in validated if a.bu != "BID"))
        extended = [_by_unit(validated, field, unit_name) if word != "Open" else "" for field, _ in EXTENDED]
        _append(ws, [req.req_id, source(req), req.text, req.quote if req.kind != "group" else "", word, answer, offered, *extended],
                fill=FILLS.get(word), bold=req.kind == "group")
        for kid in kids.get(req.req_id, []):  # sub-requirements: own RFP reference, the group's answer applies
            _append(ws, [kid.req_id, source(kid), "    " + kid.text, kid.quote, word, "", "", *[""] * len(EXTENDED)],
                    fill=FILLS.get(word), muted=True)
    _finish(ws, header_row, [w for _, w in CUSTOMER_COLUMNS])

    tr = wb.create_sheet("Tracking")
    _header(tr, 1, tracking_columns)
    for values in tracking_rows:
        _append(tr, [v[1:] if v.startswith("'") else v for v in values])  # the CSV's formula guard is not needed here
    _finish(tr, 1, TRACKING_WIDTHS)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _by_unit(validated: list, field: str, unit_name) -> str:
    """What the validated answers say in one extended field, one line per unit; blank when none of them has it (A-18)."""
    return "\n".join(f"{unit_name(a.bu)}: {v}" for a in validated if (v := getattr(a, field, None)) and str(v).strip())


def _header(ws, row: int, names: list[str]) -> None:
    for col, name in enumerate(names, 1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _append(ws, values: list, fill: str | None = None, bold: bool = False, muted: bool = False) -> None:
    r = ws.max_row + 1
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=r, column=col)
        cell.value = "" if value is None else str(value)
        cell.data_type = "s"  # always text: "=..." stays words, never a formula
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        cell.font = Font(bold=bold, color="64748B" if muted else "0F172A")
    if fill:
        ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=fill)


def _finish(ws, header_row: int, widths: list[int]) -> None:
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{max(ws.max_row, header_row)}"
